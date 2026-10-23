"""Execution logic for excel_automation skill."""

from pathlib import Path
from typing import Any
import openpyxl


def run(args: dict[str, Any]) -> dict[str, Any]:
    action = args.get("action", "read")
    filepath = args.get("filepath", "")
    sheet_name = args.get("sheet_name", "Sheet1")
    data = args.get("data", [])

    if not filepath:
        return {"status": "error", "error": "Missing filepath argument."}

    p = Path(filepath)

    if action == "create":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        for row in data:
            if isinstance(row, list):
                ws.append(row)
        p.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(p))
        return {
            "status": "ok",
            "output": f"Successfully created workbook {p.name} with {len(data)} rows.",
        }

    elif action == "read":
        if not p.exists():
            return {"status": "error", "error": f"File not found: {filepath}"}
        wb = openpyxl.load_workbook(str(p), data_only=True)
        sheets = wb.sheetnames
        ws = wb[sheet_name] if sheet_name in sheets else wb.active
        rows = list(ws.iter_rows(values_only=True))
        return {
            "status": "ok",
            "output": f"Read {len(rows)} rows from {p.name} (sheet: {ws.title}).",
            "data": rows[:50],
        }

    return {"status": "error", "error": f"Unknown action: {action}"}
