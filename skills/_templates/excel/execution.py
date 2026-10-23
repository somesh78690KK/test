"""Execution template for excel skills."""

from pathlib import Path
from typing import Any
import openpyxl


def run(args: dict[str, Any]) -> dict[str, Any]:
    filepath = args.get("filepath", "")
    p = Path(filepath)
    if not p.exists():
        wb = openpyxl.Workbook()
        wb.save(str(p))
        return {"status": "ok", "output": f"Initialized new workbook: {p.name}"}

    wb = openpyxl.load_workbook(str(p))
    return {"status": "ok", "output": f"Loaded workbook: {p.name} with sheets {wb.sheetnames}"}
